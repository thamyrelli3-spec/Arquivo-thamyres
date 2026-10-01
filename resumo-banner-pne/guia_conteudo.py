# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

AZUL = RGBColor(0x00, 0x00, 0x66)

doc = Document()
s = doc.sections[0]
for a in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
    setattr(s, a, Cm(2.0))

def p(t='', tam=11, neg=False, ital=False, cor=None, al=WD_ALIGN_PARAGRAPH.JUSTIFY,
      esp=1.15, dep=6, recuo=0):
    par = doc.add_paragraph(); par.alignment = al
    pf = par.paragraph_format; pf.line_spacing = esp
    pf.space_before = Pt(0); pf.space_after = Pt(dep)
    if recuo: pf.left_indent = Cm(recuo)
    r = par.add_run(t); r.bold = neg; r.italic = ital
    if cor is not None: r.font.color.rgb = cor
    r.font.name = 'Arial'; r.font.size = Pt(tam)
    r._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    return par

def h1(t): p(t, 15, True, False, AZUL, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 10)
def h2(t): p(t, 12.5, True, False, None, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 4)
def b(t, tam=11): p('•  ' + t, tam, False, False, None, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 4, recuo=0.6)
def nota(t): p(t, 10.5, False, True, None, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 8, recuo=0.6)

def tabela(cabecalho, linhas, larguras=None):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, c in enumerate(cabecalho):
        cel = t.rows[0].cells[i]; cel.text = ''
        r = cel.paragraphs[0].add_run(c); r.bold = True
        r.font.name = 'Arial'; r.font.size = Pt(10)
    for linha in linhas:
        cels = t.add_row().cells
        for i, c in enumerate(linha):
            cels[i].text = ''
            r = cels[i].paragraphs[0].add_run(c)
            r.font.name = 'Arial'; r.font.size = Pt(10)
    if larguras:
        for row in t.rows:
            for i, w in enumerate(larguras):
                row.cells[i].width = Cm(w)
                for par in row.cells[i].paragraphs:
                    par.paragraph_format.space_after = Pt(2)
    p('', 6, dep=10)
    return t

# ============================================================== CAPA
p('GUIA DE ESTUDO', 20, True, False, AZUL, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 2)
p('Manejo odontológico do paciente com TEA nível 3, TDAH e TAG', 13, False, False, None,
  WD_ALIGN_PARAGRAPH.CENTER, 1.15, 4)
p('III Simpósio de Pacientes com Necessidades Especiais — V JOCAM', 10.5, False, True, None,
  WD_ALIGN_PARAGRAPH.CENTER, 1.15, 14)
p('Conteúdo do banner, organizado para estudo. As dez partes vão do diagnóstico ao manejo; '
  'a Parte 11 é a revisão rápida para ler na véspera.', 11, False, False, None,
  WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 16)

# ============================================================== 1
h1('PARTE 1 — AS TRÊS CONDIÇÕES')

h2('1.1 Transtorno do Espectro Autista (TEA)')
p('Distúrbio do neurodesenvolvimento que afeta de forma persistente duas áreas: a comunicação e '
  'a interação social, e o padrão de comportamento, que se torna restrito e repetitivo. O '
  'diagnóstico é clínico, pelo DSM-5-TR, e não existe exame laboratorial que o confirme.')
p('O DSM-5 classifica em três níveis, pelo quanto de apoio a pessoa precisa:', dep=4)
tabela(['Nível', 'Denominação', 'O que significa na prática'],
 [['1', 'Exige apoio', 'Fala, mas tem dificuldade de iniciar interação; inflexibilidade atrapalha '
   'a rotina. Costuma colaborar no consultório com manejo básico.'],
  ['2', 'Exige apoio substancial', 'Fala limitada a frases curtas; comportamentos repetitivos '
   'evidentes; sofrimento claro com mudanças.'],
  ['3', 'Exige apoio muito substancial', 'Comunicação verbal mínima ou ausente; extrema '
   'dificuldade em lidar com mudança; comportamentos restritos que interferem em tudo. '
   'É o nível da nossa paciente.']],
 [1.6, 4.2, 11.2])
p('O que mais importa para a Odontologia é o quarto item dos critérios: a hiper ou '
  'hiporreatividade a estímulos sensoriais. Luz, ruído, cheiro, sabor e toque podem ser sentidos '
  'de forma muito mais intensa. É isso que transforma a cadeira odontológica — luz do refletor, '
  'ruído da alta rotação e do sugador, cheiro de eugenol, toque na boca — num ambiente hostil.')
nota('Prevalência: o CDC (relatório de 2025, ano de vigilância 2022) identificou TEA em 1 a cada '
     '31 crianças de 8 anos nos EUA, contra 1 em 36 no levantamento anterior.')

h2('1.2 Transtorno de Déficit de Atenção e Hiperatividade (TDAH)')
p('Também é transtorno do neurodesenvolvimento. Caracteriza-se por desatenção e/ou '
  'hiperatividade-impulsividade em grau que prejudica a vida da criança. Tem três apresentações:')
b('Predominantemente desatenta — não sustenta atenção, parece não ouvir, perde coisas.')
b('Predominantemente hiperativa-impulsiva — não para quieta, interrompe, age sem pensar.')
b('Combinada — as duas, é a mais comum.')
p('No consultório, o TDAH aparece como dificuldade de permanecer na cadeira, de aguardar e de '
  'seguir instrução sequencial. Em casa, aparece como dificuldade de manter rotina de escovação.')

h2('1.3 Transtorno de Ansiedade Generalizada (TAG)')
p('Preocupação excessiva e difícil de controlar, por pelo menos seis meses, acompanhada de '
  'sintomas como inquietação, irritabilidade, tensão muscular e alteração de sono. Na criança a '
  'ansiedade raramente é verbalizada: aparece como irritabilidade, queixas físicas (dor de '
  'barriga, dor de cabeça), choro, recusa e comportamento de fuga.')
p('No atendimento, o efeito é a antecipação: a criança já chega em estado de alerta esperando '
  'que algo ruim aconteça, antes mesmo de qualquer procedimento.')

h2('1.4 Por que as três juntas pioram o quadro')
p('Não é soma, é multiplicação. O TEA traz a dificuldade de comunicação e a hipersensibilidade; '
  'o TDAH impede permanecer parada e seguir instruções; o TAG faz a criança chegar já alerta. '
  'Cada um ataca um pilar diferente do manejo comportamental — e o manejo depende dos três.')
p('A coexistência é comum, não é exceção:', dep=4)
b('Metanálise de 2019 na Lancet Psychiatry: TDAH em 28% das pessoas com TEA e transtornos de '
  'ansiedade em 20%. Em amostras clínicas, como a de uma clínica de PNE, as taxas são bem maiores.')
b('Quando há os dois diagnósticos, o protocolo precisa contemplar os dois: previsibilidade e '
  'suporte visual pelo TEA, consulta curta e objetiva pelo TDAH.')

doc.add_page_break()
# ============================================================== 2
h1('PARTE 2 — O QUE ISSO CAUSA NA BOCA')
p('Ponto que a banca adora: a condição em si não causa cárie. O que causa é a cadeia de fatores '
  'que vem junto. Saiba enumerar os cinco.', dep=8)

h2('2.1 Risco aumentado de cárie e doença periodontal')
tabela(['Fator', 'Por quê'],
 [['Higiene dependente do cuidador', 'A criança não escova sozinha de forma eficaz. A qualidade '
   'da higiene é a qualidade da escovação que o cuidador consegue fazer — e ele nem sempre foi '
   'orientado.'],
  ['Hipersensibilidade oral', 'O toque da escova é aversivo. A criança recusa, o cuidador desiste, '
   'o biofilme se acumula.'],
  ['Seletividade alimentar', 'A recusa de textura leva a dieta pastosa, repetitiva e muitas vezes '
   'açucarada, com alta frequência de ingestão.'],
  ['Hipossalivação medicamentosa', 'Psicofármacos reduzem o fluxo salivar, e com ele a capacidade '
   'tampão e a autolimpeza.'],
  ['Dificuldade de acesso', 'Consultas desmarcadas, profissionais sem treinamento, diagnóstico '
   'tardio — a lesão é descoberta já extensa.']],
 [5.0, 12.0])
nota('Dado para citar: metanálise de Drumond et al. (2022) mostrou que crianças com TDAH têm '
     'razão de chances de 3,31 (IC 95% 1,25–8,73) para cárie em relação às sem o transtorno.')

h2('2.2 Os medicamentos e a boca')
tabela(['Classe / exemplo', 'Para quê', 'Repercussão bucal'],
 [['Psicoestimulantes — metilfenidato, lisdexanfetamina', 'TDAH',
   'Xerostomia, bruxismo, redução de apetite e alteração do padrão alimentar'],
  ['Antipsicóticos atípicos — risperidona, aripiprazol', 'Irritabilidade e agressividade no TEA',
   'Alteração do fluxo salivar, ganho de peso, sedação'],
  ['ISRS — sertralina, fluoxetina', 'Ansiedade',
   'Xerostomia e bruxismo']],
 [5.5, 4.0, 7.5])
p('Pergunte sempre qual medicamento a criança usa, a dose e o horário. O horário importa: uma '
  'consulta marcada no pico do estimulante pega a criança mais focada.')

h2('2.3 Outras alterações frequentes')
b('Bruxismo — comum, agravado por estimulantes e por ansiedade; leva a desgaste e sensibilidade.')
b('Traumatismo dentário — mais frequente por alterações motoras, impulsividade e crises.')
b('Autolesão — mordedura de lábio, língua e mucosa em alguns pacientes com TEA nível 3.')
b('Respiração bucal e sialorreia em parte dos casos, com repercussão em gengivite.')
b('Erosão — se houver refluxo ou ruminação, que ocorrem em parte dos pacientes.')

doc.add_page_break()
# ============================================================== 3
h1('PARTE 3 — COMO AVALIAR ANTES DE TRATAR')

h2('3.1 A anamnese ampliada')
p('É o que diferencia o atendimento. Além da anamnese comum, pergunte ao responsável:', dep=4)
b('Como a criança se comunica? Fala, aponta, usa prancha, usa aplicativo, puxa pela mão?')
b('Qual é a rotina dela? Em que horário ela está melhor?')
b('O que a incomoda? Luz forte, barulho, cheiro, toque, ser segurada, espera?')
b('O que a acalma? Objeto preferido, música, vídeo, um brinquedo específico?')
b('Como é a escovação em casa? Quem faz, quanto tempo, em que posição, com qual creme?')
b('Como é a alimentação? Quais texturas aceita, com que frequência come, o que bebe?')
b('Já foi a dentista antes? Como foi? O que deu errado?')
b('Quais medicamentos usa, em que dose e horário? Faz terapia? Com qual equipe?')
nota('Essas respostas definem o horário da consulta, quem vai atender, o que vai ser mostrado '
     'primeiro e o reforço que será usado. Não é formalidade: é o plano de tratamento.')

h2('3.2 Escala de Frankl')
p('Classificação clássica do comportamento infantil na cadeira, usada em quase todos os estudos '
  'da área. Saber de cor pega bem:', dep=4)
tabela(['Grau', 'Classificação', 'Comportamento'],
 [['1', 'Definitivamente negativo', 'Recusa o tratamento, chora intensamente, resiste fisicamente'],
  ['2', 'Negativo', 'Relutante, pouco cooperativo, retraído'],
  ['3', 'Positivo', 'Aceita o tratamento com reservas, segue orientações com cautela'],
  ['4', 'Definitivamente positivo', 'Bom vínculo, interessado, colabora plenamente']],
 [1.6, 5.4, 10.0])

h2('3.3 O cuidador é parte da equipe')
p('No paciente com TEA nível 3 o cuidador é tradutor, co-terapeuta e termômetro. Ele interpreta '
  'os sinais da criança, executa a higiene em casa e sustenta a rotina entre as consultas. A '
  'permanência dele na sala reduz a ansiedade, e orientá-lo rende mais resultado do que qualquer '
  'procedimento isolado.')

doc.add_page_break()
# ============================================================== 4
h1('PARTE 4 — MANEJO BÁSICO')
p('A diretriz da American Academy of Pediatric Dentistry (AAPD) divide o manejo em básico e '
  'avançado. O básico vem primeiro, sempre. É o conteúdo central do banner.', dep=8)

h2('4.1 Previsibilidade — o princípio que organiza tudo')
p('Para quem tem TEA, o imprevisto é a maior fonte de estresse. Então:')
b('Consultas curtas. Melhor três sessões de 15 minutos do que uma de 45.')
b('Sempre no mesmo horário, de preferência no período em que a criança está melhor.')
b('Sempre na mesma sala ou box, com a mesma equipe. Trocar o operador pode zerar o progresso.')
b('Primeira consulta apenas de adaptação, sem procedimento. Parece perda de tempo; é o que faz '
  'a segunda consulta acontecer.')
b('Avisar antes cada etapa, na mesma ordem, toda vez.')

h2('4.2 Dessensibilização progressiva')
p('Aproximação gradual e repetida do que causa aversão. A criança entra na sala, senta na '
  'cadeira, sobe e desce a cadeira, vê o espelho, encosta o espelho na mão, depois na bochecha, '
  'depois na boca. Cada etapa é repetida até ser tolerada antes de avançar para a seguinte. '
  'Programas de dessensibilização se mostraram superiores à abordagem convencional de orientação '
  'comportamental em estudos da área.')

h2('4.3 Dizer-mostrar-fazer e perguntar-dizer-perguntar')
b('Dizer-mostrar-fazer (tell-show-do): explicar em linguagem acessível, demonstrar na mão da '
  'criança ou num modelo, e só então executar. No TEA nível 3 a parte do "dizer" se apoia na '
  'imagem, não na fala.')
b('Perguntar-dizer-perguntar (ask-tell-ask): checar o que a criança entendeu antes e depois de '
  'explicar. Útil quando há alguma comunicação verbal.')

h2('4.4 Pedagogia visual')
p('O recurso com melhor evidência no TEA. Usa imagens em sequência para antecipar o que vai '
  'acontecer, substituindo a instrução verbal:')
b('Agenda visual — a sequência da consulta em figuras, que a criança acompanha e vai retirando.')
b('Pranchas de comunicação — a criança aponta o que sente ou o que quer.')
b('PECS (Picture Exchange Communication System) — sistema de comunicação por troca de figuras.')
b('TEACCH — método de ensino estruturado que organiza espaço, tempo e tarefa por pistas visuais.')
b('Histórias sociais — narrativa ilustrada da visita ao dentista, lida em casa antes.')
nota('Metanálise de Balian et al. (2021): a pedagogia visual melhora a cooperação durante o '
     'atendimento e a qualidade da escovação.')

h2('4.5 Modelagem por vídeo')
p('A criança assiste, em casa, a um vídeo curto de alguém passando pelo procedimento — de '
  'preferência outra criança, ou ela mesma em consulta anterior. Chega ao consultório sabendo o '
  'que vai acontecer.')
nota('Ensaio clínico randomizado brasileiro, de Da Silva Moro et al. (2024): a modelagem por '
     'vídeo reduziu o número de consultas necessárias para o atendimento não invasivo.')

h2('4.6 Reforço positivo')
p('Elogio descritivo e recompensa imediata a cada etapa cumprida — não só no fim. "Você abriu a '
  'boca, muito bem" vale mais que "você foi bem hoje". A recompensa deve ser algo que importa '
  'para aquela criança, definido com o cuidador na anamnese.')

h2('4.7 Ambiente sensorialmente adaptado')
p('Reduzir a carga sensorial da sala:')
b('Luz — diminuir o refletor, oferecer óculos escuros.')
b('Som — abafador de ruído ou fone com música calma; evitar conversa cruzada na sala.')
b('Cheiro — evitar eugenol e produtos de odor forte.')
b('Tato — colete com peso ou cobertor pesado, que dá sensação de contenção calmante.')
b('Visual — retirar do campo de visão os instrumentos que não serão usados.')
nota('Cermak et al. (2015), estudo piloto com 44 crianças, e Stein Duker et al. (2023), ensaio '
     'cruzado com 162 crianças autistas: o ambiente adaptado reduziu o estresse fisiológico, '
     'medido no corpo da criança, em todas as fases da consulta, além do sofrimento comportamental.')

doc.add_page_break()
# ============================================================== 5
h1('PARTE 5 — MANEJO AVANÇADO')
p('Entra quando o básico não é suficiente. Nunca é a primeira escolha, e sempre exige '
  'consentimento informado e registro em prontuário.', dep=8)

h2('5.1 Estabilização protetora')
p('Restrição do movimento para proteger a criança e a equipe durante o procedimento. Regras:')
b('Só quando necessário, nunca por conveniência ou pressa.')
b('Sempre a técnica menos restritiva que permita atender com segurança.')
b('Com consentimento informado dos riscos, benefícios e alternativas, inclusive a de não tratar '
  'ou adiar.')
b('Registrada em prontuário: indicação, tipo, duração e resposta do paciente.')
b('Nunca como punição, e nunca sem supervisão contínua.')

h2('5.2 Sedação')
p('Vai da sedação mínima com óxido nitroso e oxigênio até a sedação moderada e profunda. Exige '
  'profissional habilitado, monitorização e estrutura de emergência. Útil quando o procedimento é '
  'necessário, a criança não tolera, mas não se justifica a anestesia geral.')

h2('5.3 Anestesia geral')
p('Em ambiente hospitalar. Indicada quando há grande necessidade de tratamento, impossibilidade '
  'de manejo e risco de o adiamento agravar o quadro. Resolve a demanda acumulada numa sessão, '
  'mas não ensina a criança a aceitar o atendimento — por isso o condicionamento continua depois.')

h2('5.4 A escada — como responder se perguntarem a ordem')
p('1. Manejo básico e prevenção. 2. Adaptação do ambiente. 3. Dessensibilização em várias '
  'sessões. 4. Sedação mínima. 5. Estabilização protetora, se indicada. 6. Sedação moderada ou '
  'profunda. 7. Anestesia geral. Sobe-se um degrau por vez, e só quando o anterior falhou.')

doc.add_page_break()
# ============================================================== 6
h1('PARTE 6 — PREVENÇÃO: O QUE REALMENTE MUDA O DESFECHO')
p('Quanto mais prevenção, menos necessidade dos degraus de cima. É a tese da conclusão do '
  'banner.', dep=8)
b('Escovação supervisionada pelo cuidador, com técnica orientada e posição definida — a posição '
  'joelho a joelho, ou a criança deitada com a cabeça no colo do cuidador, funciona bem.')
b('Escova elétrica — útil quando a higiene depende de terceiros; em alguns pacientes a vibração '
  'é aversiva, então teste antes.')
b('Creme dental fluoretado em quantidade adequada à idade; se a criança não cospe, usar quantidade '
  'reduzida.')
b('Verniz fluoretado em aplicações periódicas — rápido, bem tolerado e de alto impacto.')
b('Selantes nos molares, quando houver cooperação.')
b('Diamino fluoreto de prata — paralisa a lesão de cárie sem preparo cavitário, sem ruído e sem '
  'anestesia. Escurece a lesão, o que precisa ser explicado e consentido.')
b('Tratamento restaurador atraumático (ART) — remoção com instrumento manual e restauração com '
  'ionômero, sem alta rotação.')
b('Orientação dietética focada em frequência, não só em quantidade.')
b('Retornos curtos e frequentes, que mantêm o vínculo e o condicionamento.')

doc.add_page_break()
# ============================================================== 7
h1('PARTE 7 — DIREITOS E LEGISLAÇÃO')
tabela(['Norma', 'O que garante'],
 [['Lei 12.764/2012 — Política Nacional de Proteção dos Direitos da Pessoa com TEA',
   'Considera a pessoa com TEA como pessoa com deficiência para todos os efeitos legais, com '
   'direito a atendimento integral em saúde.'],
  ['Lei 13.977/2020 — Lei Romeo Mion',
   'Cria a Carteira de Identificação da Pessoa com TEA (CIPTEA), que dá acesso a atendimento '
   'prioritário.'],
  ['Lei 13.146/2015 — Lei Brasileira de Inclusão',
   'Define os tipos de barreira (urbanística, arquitetônica, de transporte, de comunicação, '
   'atitudinal e tecnológica) que devem ser removidas.'],
  ['Guia de Atenção à Saúde Bucal da Pessoa com Deficiência (MS, 2019)',
   'Orienta a rede SUS sobre manejo e cuidado, com capítulo sobre TEA.']],
 [6.0, 11.0])

# ============================================================== 8
h1('PARTE 8 — AS EVIDÊNCIAS DO BANNER NUMA TABELA')
tabela(['Referência', 'Tipo de estudo', 'O que mostrou'],
 [['AAPD (2024)', 'Diretriz', 'Separa manejo básico de avançado e define quando escalar'],
  ['Balian et al. (2021)', 'Revisão sistemática com metanálise', 'Pedagogia visual melhora '
   'cooperação e higiene bucal'],
  ['Stein Duker et al. (2023)', 'ECR cruzado, 162 crianças', 'Ambiente adaptado reduz estresse '
   'fisiológico e comportamental'],
  ['Cermak et al. (2015)', 'ECR piloto, 44 crianças', 'Mesma linha, primeiro estudo do tema'],
  ['Da Silva Moro et al. (2024)', 'ECR, autores brasileiros', 'Modelagem por vídeo reduz o número '
   'de consultas'],
  ['Drumond et al. (2022)', 'Metanálise', 'Cárie em TDAH: OR 3,31 (IC 95% 1,25–8,73)'],
  ['AlBhaisi et al. (2022)', 'Revisão sistemática', 'Técnicas psicológicas funcionam, mas a '
   'evidência é inconclusiva quanto à força'],
  ['Bezerra et al. (2023)', 'Revisão de literatura, em português', 'Panorama nacional do '
   'atendimento à criança com TEA']],
 [4.6, 4.6, 7.8])
nota('A referência mais forte do conjunto é Stein Duker (2023). A mais fácil de ler inteira é '
     'Bezerra (2023), em português e gratuita. Leia essa primeiro.')

doc.add_page_break()
# ============================================================== 9
h1('PARTE 9 — PERGUNTAS DE BANCA, COM RESPOSTA')
qa = [
 ('Por que não sedar logo, já que a paciente não colabora?',
  'Porque a AAPD coloca sedação e anestesia geral como manejo avançado, para depois da falha do '
  'básico. Além do risco e do custo, a sedação resolve a consulta mas não ensina a criança a '
  'aceitar o atendimento. O condicionamento tem efeito duradouro.'),
 ('Por que a criança com TEA tem mais cárie?',
  'Não é a condição em si. É a soma de higiene dependente do cuidador, hipersensibilidade oral '
  'que dificulta a escovação, dieta seletiva e cariogênica e hipossalivação causada pelos '
  'psicofármacos.'),
 ('O que muda no atendimento por ser nível 3 e não nível 1?',
  'No nível 3 a comunicação verbal é mínima, então todo o manejo se apoia em recurso visual e em '
  'rotina, não em explicação falada. E a alteração sensorial costuma ser mais intensa, o que torna '
  'a adaptação do ambiente indispensável, não opcional.'),
 ('Qual a diferença entre estabilização protetora e imobilização?',
  'Estabilização protetora é indicada, consentida, registrada em prontuário e feita na técnica '
  'menos restritiva possível. Imobilização sem indicação e sem consentimento é conduta '
  'inadequada.'),
 ('A pedagogia visual funciona mesmo ou é teoria?',
  'A metanálise de Balian mostra melhora de cooperação e de higiene bucal. A limitação é que os '
  'estudos são pequenos e heterogêneos, como o próprio AlBhaisi aponta.'),
 ('Como vocês lidariam com a hipersensibilidade na hora da escovação?',
  'Dessensibilização gradual do toque, começando fora da boca; escova de cerdas macias e cabeça '
  'pequena; testar escova elétrica, que ajuda em parte dos casos e piora em outros; posição joelho '
  'a joelho; e creme sem sabor forte, porque o sabor é um gatilho frequente.'),
 ('O TDAH muda alguma coisa no manejo, ou é só o TEA?',
  'Muda. O TDAH exige consulta curta e objetiva, com poucas instruções por vez, e aumenta o risco '
  'de cárie por falha de rotina de higiene e por impulsividade alimentar. É ele que justifica '
  'sessões de 15 minutos em vez de 45.'),
 ('E a ansiedade, como vocês abordam?',
  'Com previsibilidade. A ansiedade generalizada faz a criança antecipar o pior, então saber '
  'exatamente o que vai acontecer, na mesma ordem, com a mesma equipe, é o que reduz a '
  'antecipação. Vídeo e agenda visual em casa atacam exatamente isso.'),
 ('Qual recurso vocês implantariam amanhã na clínica?',
  'Agenda visual impressa e vídeo gravado com o celular. Custo quase zero, e são justamente os '
  'dois com ensaio clínico mostrando ganho de cooperação e redução do número de consultas.'),
 ('Se nada funcionar, o que vocês fazem?',
  'Mantemos a adequação do meio bucal e o controle preventivo, e escalamos conforme a necessidade '
  'e a urgência: sedação e, se indicado, anestesia geral, sempre com consentimento e em conjunto '
  'com a família e com o médico que acompanha a criança.'),
 ('Vocês citam que não existe protocolo único. Então o que o trabalho conclui?',
  'Que a conduta é individualizada, mas os princípios são constantes: anamnese ampliada, '
  'previsibilidade, suporte visual, adaptação sensorial, reforço positivo e prevenção. O que varia '
  'é a combinação e o ritmo, não os princípios.'),
 ('Por que a primeira consulta sem procedimento não é perda de tempo?',
  'Porque o custo de uma primeira consulta malsucedida é a recusa nas seguintes. A consulta de '
  'adaptação constrói o vínculo que permite todas as outras — é investimento, não desperdício.'),
]
for q, a in qa:
    p(q, 11, True, False, None, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 0)
    p(a, 11, False, False, None, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 8, recuo=0.6)

doc.add_page_break()
# ============================================================== 10
h1('PARTE 10 — GLOSSÁRIO')
tabela(['Termo', 'Significado'],
 [['DSM-5-TR', 'Manual Diagnóstico e Estatístico de Transtornos Mentais, 5ª edição, texto revisado'],
  ['Dessensibilização', 'Exposição gradual e repetida ao estímulo aversivo até que seja tolerado'],
  ['Dizer-mostrar-fazer', 'Explicar, demonstrar e só então executar'],
  ['Pedagogia visual', 'Uso de imagens em sequência para antecipar e organizar o procedimento'],
  ['PECS', 'Sistema de comunicação por troca de figuras'],
  ['TEACCH', 'Método de ensino estruturado por pistas visuais de espaço, tempo e tarefa'],
  ['Modelagem por vídeo', 'Assistir a alguém passando pelo procedimento antes de vivê-lo'],
  ['Escala de Frankl', 'Classificação do comportamento infantil na cadeira, de 1 a 4'],
  ['Estabilização protetora', 'Restrição de movimento indicada, consentida e registrada'],
  ['ART', 'Tratamento restaurador atraumático, sem alta rotação'],
  ['Diamino fluoreto de prata', 'Agente que paralisa a lesão de cárie sem preparo cavitário'],
  ['Xerostomia', 'Sensação de boca seca, geralmente por redução do fluxo salivar'],
  ['CIPTEA', 'Carteira de Identificação da Pessoa com Transtorno do Espectro Autista']],
 [5.0, 12.0])

# ============================================================== 11
h1('PARTE 11 — REVISÃO DE DEZ MINUTOS')
p('Para ler antes de apresentar. Se você souber estes dez pontos, sustenta qualquer conversa '
  'sobre o banner.', dep=8)
for i, t in enumerate([
 'TEA nível 3 é o que exige apoio muito substancial: comunicação verbal mínima, comportamento '
 'repetitivo e alteração sensorial.',
 'A alteração sensorial é o que mais importa na Odontologia: luz, ruído, cheiro e toque são '
 'sentidos de forma intensificada.',
 'TDAH impede permanecer parada e seguir instrução; TAG faz a criança chegar já em alerta. Os '
 'três juntos derrubam a colaboração.',
 'O risco de cárie não vem da condição: vem de higiene dependente do cuidador, hipersensibilidade '
 'oral, dieta seletiva e boca seca por medicamento.',
 'Metanálise de 2022: crianças com TDAH têm OR 3,31 para cárie.',
 'A anamnese ampliada é o plano de tratamento: rotina, comunicação, gatilhos, reforços, '
 'medicamentos.',
 'Previsibilidade é o princípio central: consulta curta, mesmo horário, mesma sala, mesma equipe, '
 'primeira sessão só de adaptação.',
 'O condicionamento é gradual: dessensibilização, dizer-mostrar-fazer, pedagogia visual, vídeo e '
 'reforço positivo a cada etapa.',
 'O ambiente adaptado reduz estresse fisiológico comprovadamente — ECR com 162 crianças, 2023.',
 'Contenção, sedação e anestesia geral só depois da falha do básico, com consentimento. E a '
 'prevenção é o que evita chegar lá.',
]):
    p('%2d.  %s' % (i + 1, t), 11, False, False, None, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 5, recuo=0.8)

doc.save('Guia_Estudo_Conteudo_TEA_TDAH_TAG.docx')
print('ok')
