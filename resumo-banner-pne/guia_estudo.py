# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
sec = doc.sections[0]
for a in ('top_margin','bottom_margin','left_margin','right_margin'):
    setattr(sec, a, Cm(2.0))

def p(texto='', tam=11, neg=False, al=WD_ALIGN_PARAGRAPH.JUSTIFY, esp=1.15, dep=6, recuo=0):
    par = doc.add_paragraph(); par.alignment = al
    pf = par.paragraph_format; pf.line_spacing = esp
    pf.space_before = Pt(0); pf.space_after = Pt(dep)
    if recuo: pf.left_indent = Cm(recuo)
    r = par.add_run(texto); r.bold = neg
    r.font.name = 'Arial'; r.font.size = Pt(tam)
    r._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    return par

def h(texto, tam=13):
    p(texto, tam, True, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 8)

def bullet(texto, tam=11):
    p('•  ' + texto, tam, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 4, recuo=0.6)

p('GUIA DE ESTUDO — BANNER DE PNE', 14, True, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 4)
p('Manejo odontológico do paciente com TEA nível 3, TDAH e TAG', 12, False,
  WD_ALIGN_PARAGRAPH.CENTER, 1.15, 16)

h('1. DE ONDE SAIU CADA PARTE DO RESUMO')
p('Cada trecho do resumo está sustentado por uma das referências. É isso que você responde '
  'se a banca perguntar "de onde você tirou isso?".', dep=8)

mapa = [
 ('"TEA nível 3 é o que exige maior suporte"',
  'Classificação do DSM-5-TR por nível de suporte (1, 2 e 3). O nível 3 é descrito como '
  '"exigindo apoio muito substancial". Base conceitual, aparece também em BEZERRA et al. (2023).'),
 ('"maior risco de cárie e doença periodontal"',
  'DRUMOND et al. (2022) — metanálise em crianças com TDAH. BEZERRA et al. (2023) para o TEA.'),
 ('"higiene depende do cuidador, dieta seletiva, psicofármacos reduzem o fluxo salivar"',
  'BEZERRA et al. (2023) e a literatura nacional sobre TEA; a hipossalivação é efeito adverso '
  'conhecido de metilfenidato, risperidona e antidepressivos.'),
 ('"anamnese detalhada com o responsável"',
  'AAPD (2024) — a diretriz exige registrar histórico médico, temperamento, comportamento em '
  'atendimentos anteriores e técnicas já utilizadas antes de escolher o manejo.'),
 ('"consultas curtas, mesmo horário, mesma sala, mesma equipe"',
  'BEZERRA et al. (2023) e AAPD (2024): previsibilidade e redução de estímulos.'),
 ('"dizer-mostrar-fazer e reforço positivo"',
  'AAPD (2024) — as duas estão na lista de manejo básico (basic behavior guidance).'),
 ('"dessensibilização"',
  'AAPD (2024) e ALBHAISI et al. (2022): aproximação gradual e repetida do procedimento.'),
 ('"pranchas e agendas de figuras" (pedagogia visual)',
  'BALIAN et al. (2021) — revisão sistemática com metanálise; é a referência mais forte do resumo '
  'para esse ponto.'),
 ('"vídeos que mostram o procedimento antes da consulta"',
  'ALBHAISI et al. (2022) — video modeling entre as técnicas psicológicas revisadas.'),
 ('"diminuir luz, ruído e odores"',
  'CERMAK et al. (2015) e STEIN DUKER et al. (2023) — ambiente sensorialmente adaptado.'),
 ('"contenção protetora, sedação e anestesia geral reservadas"',
  'AAPD (2024) — são classificadas como manejo avançado, usadas quando o básico não basta e '
  'sempre com consentimento informado.'),
 ('"não existe um protocolo único"',
  'ALBHAISI et al. (2022) — a própria revisão conclui que a evidência ainda é inconclusiva e '
  'heterogênea, o que impede um protocolo universal.'),
]
for frase, fonte in mapa:
    p(frase, 11, True, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 0)
    p(fonte, 11, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 8, recuo=0.6)

doc.add_page_break()
h('2. RESUMO DAS REFERÊNCIAS')

refs = [
 ('AAPD (2024) — Behavior guidance for the pediatric dental patient',
  'Reference Manual of Pediatric Dentistry, p. 358-378. Diretriz da Academia Americana de '
  'Odontopediatria, revisada periodicamente.',
  ['O que é: não é pesquisa, é diretriz. Serve para justificar conduta.',
   'Divide o manejo em básico e avançado. Básico: comunicação, imagens positivas antes da consulta, '
   'observação de outro paciente, dizer-mostrar-fazer, perguntar-dizer-perguntar, controle de voz, '
   'comunicação não verbal, reforço positivo, distração e dessensibilização.',
   'Avançado: contenção protetora, sedação e anestesia geral.',
   'Manda documentar histórico médico, temperamento, consentimento informado com riscos e '
   'alternativas, avaliação de dor, urgência do tratamento e comportamento em consultas anteriores.',
   'Sobre contenção: usar apenas quando necessário para proteger a criança e a equipe, e sempre a '
   'técnica menos restritiva possível (há uma diretriz específica só sobre isso).',
   'Por que usei: sustenta quase toda a parte de conduta do resumo.']),

 ('BALIAN et al. (2021) — Pedagogia visual',
  'Int J Environ Res Public Health, v. 18, n. 2, p. 789. Revisão sistemática com metanálise.',
  ['Pergunta: a pedagogia visual melhora a cooperação e a higiene bucal em crianças com TEA?',
   'Pedagogia visual = pranchas, sequências de figuras, agenda visual, histórias sociais, cartões '
   'que mostram cada etapa do que vai acontecer.',
   'Resultado: melhora as habilidades de escovação e a higiene bucal, e aumenta o nível de '
   'cooperação durante o atendimento. Material visual adaptado culturalmente reduziu a ansiedade '
   'em comparação ao grupo controle.',
   'Por que usei: é a evidência mais forte a favor das pranchas e agendas de figuras.']),

 ('ALBHAISI et al. (2022) — Técnicas psicológicas',
  'BMC Oral Health, v. 22. Revisão sistemática de literatura.',
  ['Revisa as técnicas não farmacológicas usadas com crianças com TEA: dizer-mostrar-fazer, '
   'dessensibilização, modelagem por vídeo, reforço positivo, suportes visuais.',
   'Resultado: os estudos mostram efeitos favoráveis, mas a evidência é considerada inconclusiva '
   'quanto à força, por causa da heterogeneidade metodológica e das amostras pequenas.',
   'Por que usei: sustenta a frase da conclusão sobre não existir protocolo único. É um ponto '
   'honesto do trabalho, e a banca costuma gostar quando o aluno reconhece a limitação.']),

 ('CERMAK et al. (2015) — Ambiente sensorialmente adaptado (estudo piloto)',
  'J Autism Dev Disord, v. 45, n. 9, p. 2876-2888. Ensaio clínico randomizado piloto.',
  ['44 crianças de 6 a 12 anos: 22 com TEA e 22 com desenvolvimento típico.',
   'Cada criança passou por duas profilaxias, uma em ambiente convencional e outra em ambiente '
   'adaptado, em ordem sorteada, com 3 a 4 meses de intervalo.',
   'Adaptações: luz reduzida, projeção em movimento no teto (peixes, bolhas), música calma e colete '
   'com peso que envolve a criança.',
   'Resultado: menos estresse, menos desconforto sensorial e menor percepção de dor no ambiente '
   'adaptado.',
   'Por que usei: é a origem do conceito de adaptação sensorial do consultório.']),

 ('STEIN DUKER et al. (2023) — Ambiente sensorialmente adaptado (estudo grande)',
  'JAMA Network Open, v. 6, n. 6, e2316346. Ensaio clínico randomizado cruzado.',
  ['É a continuação do estudo do Cermak, agora com 162 crianças autistas.',
   'Resultado: estresse fisiológico significativamente menor no ambiente adaptado em todas as fases '
   'da consulta, indicando menor ativação simpática, ou seja, a criança mais relaxada.',
   'Também reduziu o sofrimento comportamental, e o método foi considerado seguro.',
   'Por que usei: é a evidência de melhor qualidade do resumo. Se a banca perguntar qual o estudo '
   'mais forte que você citou, é este.']),

 ('DRUMOND et al. (2022) — Cárie e TDAH',
  'Caries Research, v. 56, n. 1, p. 3-14. Revisão sistemática com metanálise.',
  ['Compara a ocorrência de cárie em crianças com TDAH e sem TDAH.',
   'Resultado: crianças com TDAH têm mais chance de apresentar cárie (OR = 3,31; IC 95% 1,25-8,73).',
   'Explicações apontadas: dificuldade de manter rotina de higiene, falta de motivação, maior '
   'acúmulo de biofilme, sangramento gengival e consumo de alimentos cariogênicos por impulsividade.',
   'Por que usei: sustenta a frase sobre maior risco de cárie, na parte do TDAH.']),

 ('BEZERRA; ASSIS; SANTOS (2023) — Atendimento odontológico de crianças com TEA',
  'Brazilian Journal of Health Review, v. 6, n. 3, p. 13155-13171. Revisão de literatura, em '
  'português e com acesso livre.',
  ['Revisão nacional sobre as características das crianças com TEA e as estratégias de atendimento.',
   'Aborda as alterações bucais mais comuns, a dificuldade de higiene, a importância do cuidador e '
   'as técnicas de manejo.',
   'Por que usei: é a referência em português, a mais fácil de ler inteira e a melhor para você '
   'estudar primeiro. Leia esta antes das outras.']),
]
for titulo, dados, itens in refs:
    p(titulo, 11.5, True, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 0)
    p(dados, 10.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 4)
    for i in itens:
        bullet(i)
    p('', 8, dep=6)

doc.add_page_break()
h('3. O QUE ESTUDAR ALÉM DAS REFERÊNCIAS')
for t in [
 'DSM-5-TR: os três níveis de suporte do TEA e o que muda entre eles. Você precisa saber explicar '
 'por que o nível 3 é o mais grave.',
 'TDAH: os três tipos (desatento, hiperativo-impulsivo e combinado) e os medicamentos mais usados '
 '(metilfenidato, lisdexanfetamina) com seus efeitos na boca.',
 'TAG: como a ansiedade aparece na criança (irritabilidade, queixas físicas, recusa) e por que ela '
 'piora a aceitação do atendimento.',
 'Medicamentos e boca seca: metilfenidato, risperidona, aripiprazol e antidepressivos. Saiba dizer '
 'qual o efeito bucal de cada um.',
 'Seletividade alimentar no TEA: por que a dieta costuma ser pastosa, doce e repetitiva.',
 'Bruxismo, automutilação e traumatismo dentário nesses pacientes.',
 'TEACCH e PECS: são os dois métodos de comunicação estruturada citados na literatura de TEA. '
 'Vale saber o que cada sigla significa.',
 'Lei 12.764/2012 (Política Nacional de Proteção dos Direitos da Pessoa com TEA, que garante à '
 'pessoa com autismo os mesmos direitos da pessoa com deficiência) e Lei 13.977/2020, a Lei Romeo '
 'Mion, que criou a carteira de identificação (CIPTEA).',
 'Prevalência: dados do CDC de 2025 (ano de vigilância 2022) apontam 1 em cada 31 crianças de 8 '
 'anos identificadas com TEA nos EUA, contra 1 em 36 no levantamento anterior. Bom número de '
 'abertura para a apresentação.',
 'Prevenção: verniz fluoretado, selantes, escovação supervisionada pelo cuidador e ART. É o que '
 'você vai defender como prioridade.',
]:
    bullet(t)

p('', 8, dep=10)
h('4. PERGUNTAS QUE PODEM TE FAZER')
qa = [
 ('Por que não sedar logo, já que a paciente não colabora?',
  'Porque a AAPD coloca sedação e anestesia geral como manejo avançado, indicado quando o básico '
  'falha. Além do risco e do custo, a sedação resolve a consulta, mas não ensina a criança a '
  'aceitar o atendimento. O condicionamento tem efeito duradouro.'),
 ('Qual a diferença entre contenção protetora e imobilização?',
  'Contenção protetora é a estabilização, com consentimento e registro em prontuário, na técnica '
  'menos restritiva que permita atender com segurança. Imobilização sem consentimento e sem '
  'indicação é conduta inadequada.'),
 ('A pedagogia visual funciona mesmo ou é só teoria?',
  'Funciona: a metanálise do Balian mostra melhora da cooperação e da higiene bucal. A limitação é '
  'que os estudos são pequenos e heterogêneos.'),
 ('Por que a criança com TEA tem mais cárie?',
  'Não é a condição em si. É a soma de higiene dependente do cuidador, dieta seletiva e cariogênica, '
  'boca seca por medicamento e dificuldade de escovar por hipersensibilidade na boca.'),
 ('O que você faria se nada disso funcionasse?',
  'Manteria a adequação do meio bucal e o controle preventivo, escalaria para sedação ou anestesia '
  'geral conforme a necessidade e a urgência do tratamento, sempre com consentimento e em conjunto '
  'com a família e o médico que acompanha a criança.'),
]
for q, a in qa:
    p(q, 11, True, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 0)
    p(a, 11, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 8, recuo=0.6)

p('', 8, dep=10)
h('5. ONDE LER CADA UMA')
links = [
 'BEZERRA et al. (2023): ojs.brazilianjournals.com.br/ojs/index.php/BJHR/article/view/60794',
 'AAPD (2024): aapd.org/globalassets/media/policies_guidelines/bp_behavguide.pdf',
 'AAPD, contenção protetora: aapd.org/media/Policies_Guidelines/BP_Protective.pdf',
 'BALIAN et al. (2021): ncbi.nlm.nih.gov/pmc/articles/PMC7832292',
 'ALBHAISI et al. (2022): bmcoralhealth.biomedcentral.com/articles/10.1186/s12903-022-02200-7',
 'CERMAK et al. (2015): pubmed.ncbi.nlm.nih.gov/25931290',
 'STEIN DUKER et al. (2023): ncbi.nlm.nih.gov/pmc/articles/PMC10238943',
 'DRUMOND et al. (2022): karger.com/cre/article/56/1/3',
]
for l in links:
    bullet(l, 10.5)

doc.save('Guia_de_estudo_Banner_PNE.docx')
print('ok')
