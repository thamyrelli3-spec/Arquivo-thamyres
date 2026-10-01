# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

AZUL = RGBColor(0x00, 0x00, 0x66)
doc = Document()
s = doc.sections[0]
for a in ('top_margin', 'bottom_margin'): setattr(s, a, Cm(1.5))
for a in ('left_margin', 'right_margin'): setattr(s, a, Cm(1.8))

def p(t='', tam=9.5, neg=False, ital=False, cor=None, al=WD_ALIGN_PARAGRAPH.JUSTIFY,
      esp=1.0, dep=3, recuo=0):
    par = doc.add_paragraph(); par.alignment = al
    pf = par.paragraph_format; pf.line_spacing = esp
    pf.space_before = Pt(0); pf.space_after = Pt(dep)
    if recuo: pf.left_indent = Cm(recuo)
    r = par.add_run(t); r.bold = neg; r.italic = ital
    if cor is not None: r.font.color.rgb = cor
    r.font.name = 'Arial'; r.font.size = Pt(tam)
    r._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    return par

def h(t): p(t, 12, True, False, AZUL, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 5)
def sub(t): p(t, 10, True, False, None, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 2)
def b(t): p('•  ' + t, 9.5, False, False, None, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.0, 2, recuo=0.45)

def tab(cab, linhas, larg):
    t = doc.add_table(rows=1, cols=len(cab)); t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, c in enumerate(cab):
        cel = t.rows[0].cells[i]; cel.text = ''
        r = cel.paragraphs[0].add_run(c); r.bold = True
        r.font.name = 'Arial'; r.font.size = Pt(9)
    for linha in linhas:
        cels = t.add_row().cells
        for i, c in enumerate(linha):
            cels[i].text = ''
            r = cels[i].paragraphs[0].add_run(c)
            r.font.name = 'Arial'; r.font.size = Pt(9)
    for row in t.rows:
        for i, w in enumerate(larg):
            row.cells[i].width = Cm(w)
            for par in row.cells[i].paragraphs:
                par.paragraph_format.space_after = Pt(1)
                par.paragraph_format.line_spacing = 1.0
    p('', 5, dep=6)

# ---------------------------------------------------------------- cabeçalho
p('GUIA RÁPIDO DE ESTUDO', 15, True, False, AZUL, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 1)
p('Manejo odontológico do paciente com TEA nível 3, TDAH e TAG', 10.5, False, False, None,
  WD_ALIGN_PARAGRAPH.CENTER, 1.0, 8)

# ---------------------------------------------------------------- 1
h('1. AS TRÊS CONDIÇÕES')
tab(['Condição', 'O que é', 'O que atrapalha no consultório'],
 [['TEA nível 3', 'Transtorno do neurodesenvolvimento que afeta comunicação, interação social e '
   'comportamento. Nível 3 = exige apoio muito substancial (DSM-5).',
   'Comunicação verbal mínima e hiper-reatividade sensorial: luz, ruído, cheiro e toque são '
   'sentidos de forma intensificada.'],
  ['TDAH', 'Desatenção e/ou hiperatividade-impulsividade. Três apresentações: desatenta, '
   'hiperativa-impulsiva e combinada.',
   'Não permanece parada na cadeira, não aguarda, não segue instrução sequencial. Em casa, '
   'falha a rotina de escovação.'],
  ['TAG', 'Preocupação excessiva por 6 meses ou mais. Na criança aparece como irritabilidade, '
   'queixa física, choro e recusa.',
   'Antecipação: a criança já chega em estado de alerta, antes de qualquer procedimento.']],
 [2.4, 7.0, 8.0])
p('Cada transtorno derruba um pilar diferente do manejo — e o manejo depende dos três. A '
  'coexistência é comum: metanálise da Lancet Psychiatry (2019) encontrou TDAH em 28% e '
  'transtornos de ansiedade em 20% das pessoas com TEA; em amostras clínicas, muito mais.')

# ---------------------------------------------------------------- 2
h('2. O QUE ISSO CAUSA NA BOCA')
p('A condição em si não causa cárie. Causa a cadeia de fatores que vem junto — saiba enumerar:',
  dep=3)
b('Higiene dependente do cuidador, que nem sempre foi orientado.')
b('Hipersensibilidade oral: o toque da escova é aversivo, a criança recusa, o biofilme acumula.')
b('Seletividade alimentar: dieta pastosa, repetitiva e açucarada, com alta frequência.')
b('Hipossalivação medicamentosa: estimulantes (metilfenidato), antipsicóticos (risperidona) e '
  'ISRS (sertralina) reduzem o fluxo salivar; estimulantes e ISRS também causam bruxismo.')
b('Acesso difícil: consultas desmarcadas, equipe sem treinamento, diagnóstico tardio.')
p('Também são frequentes bruxismo, traumatismo dentário e, no TEA nível 3, autolesão por '
  'mordedura de lábio e mucosa.')

# ---------------------------------------------------------------- 3
h('3. COMO CONDUZIR O ATENDIMENTO')
sub('Antes: anamnese ampliada com o responsável')
p('Como a criança se comunica · qual a rotina e o melhor horário · o que a incomoda · o que a '
  'acalma · como é a escovação e a alimentação em casa · experiências anteriores no dentista · '
  'medicamentos, doses e horários. Essas respostas são o plano de tratamento, não formalidade.')
sub('Previsibilidade — o princípio que organiza tudo')
b('Consultas curtas: três de 15 minutos valem mais que uma de 45.')
b('Mesmo horário, mesma sala, mesma equipe. Trocar o operador pode zerar o progresso.')
b('Primeira consulta só de adaptação, sem procedimento. É o que faz a segunda acontecer.')
sub('Condicionamento gradual')
b('Dessensibilização: aproximação repetida e progressiva — entrar na sala, sentar, ver o espelho, '
  'encostar na mão, na bochecha, na boca. Só avança quando a etapa anterior é tolerada.')
b('Dizer-mostrar-fazer: explicar, demonstrar, executar. No nível 3 o "dizer" se apoia na imagem.')
b('Pedagogia visual: agenda visual da consulta em figuras, pranchas de comunicação, PECS, TEACCH '
  'e histórias sociais. É o recurso com melhor evidência.')
b('Modelagem por vídeo: a criança assiste em casa a alguém passando pelo procedimento.')
b('Reforço positivo imediato a cada etapa cumprida, com algo que importa para aquela criança.')
sub('Ambiente sensorialmente adaptado')
p('Reduzir luz do refletor e oferecer óculos escuros · abafador de ruído ou música calma · evitar '
  'eugenol e odores fortes · colete ou cobertor com peso · tirar do campo de visão o instrumental '
  'que não será usado · manter o cuidador na sala.')

# ---------------------------------------------------------------- 4
h('4. QUANDO ESCALAR, E A ORDEM')
p('A AAPD separa manejo básico de avançado. O avançado só entra depois da falha do básico, sempre '
  'com consentimento informado e registro em prontuário.', dep=3)
p('A escada: 1) manejo básico e prevenção → 2) adaptação do ambiente → 3) dessensibilização em '
  'várias sessões → 4) óxido nitroso → 5) estabilização protetora, se indicada → 6) sedação '
  'moderada ou profunda → 7) anestesia geral. Sobe-se um degrau por vez.', dep=3)
p('Sobre a estabilização protetora: só quando necessário para proteger a criança e a equipe, '
  'sempre na técnica menos restritiva, com consentimento e registro. Nunca por pressa e nunca '
  'como punição.')

# ---------------------------------------------------------------- 5
h('5. PREVENÇÃO — O QUE MUDA O DESFECHO')
p('Escovação supervisionada pelo cuidador, com técnica e posição orientadas (joelho a joelho) · '
  'escova elétrica, testando a tolerância à vibração · creme fluoretado em quantidade adequada · '
  'verniz fluoretado periódico · selantes · diamino fluoreto de prata, que paralisa a lesão sem '
  'preparo cavitário (escurece, precisa ser consentido) · ART, sem alta rotação · orientação '
  'dietética focada em frequência · retornos curtos e frequentes.')

# ---------------------------------------------------------------- 6
h('6. NÚMEROS E ESTUDOS PARA CITAR')
tab(['Dado', 'Fonte'],
 [['Crianças com TDAH têm OR 3,31 (IC 95% 1,25–8,73) para cárie', 'Drumond et al., 2022 — metanálise'],
  ['Ambiente adaptado reduz estresse fisiológico e comportamental em todas as fases da consulta',
   'Stein Duker et al., 2023 — ECR cruzado, 162 crianças'],
  ['Pedagogia visual melhora cooperação e qualidade da escovação', 'Balian et al., 2021 — metanálise'],
  ['Modelagem por vídeo reduz o número de consultas necessárias',
   'Da Silva Moro et al., 2024 — ECR brasileiro'],
  ['Evidência das técnicas é favorável, mas inconclusiva quanto à força',
   'AlBhaisi et al., 2022 — revisão sistemática'],
  ['Manejo básico x avançado e critérios para escalar', 'AAPD, 2024 — diretriz'],
  ['TEA em 1 a cada 31 crianças de 8 anos', 'CDC, 2025 (ano de vigilância 2022)']],
 [9.0, 8.4])

# ---------------------------------------------------------------- 7
h('7. SEIS PERGUNTAS PROVÁVEIS')
for q, a in [
 ('Por que não sedar logo, já que a paciente não colabora?',
  'A AAPD coloca sedação e anestesia geral como manejo avançado, depois da falha do básico. A '
  'sedação resolve a consulta, mas não ensina a criança a aceitar o atendimento.'),
 ('Por que a criança com TEA tem mais cárie?',
  'Não é a condição: é higiene dependente do cuidador, hipersensibilidade oral, dieta seletiva e '
  'boca seca por medicamento.'),
 ('O que muda por ser nível 3 e não nível 1?',
  'A comunicação verbal é mínima, então o manejo se apoia em recurso visual e rotina, não em '
  'explicação falada. E a alteração sensorial é mais intensa, o que torna a adaptação do ambiente '
  'indispensável.'),
 ('Estabilização protetora é o mesmo que imobilização?',
  'Não. A estabilização é indicada, consentida, registrada e na técnica menos restritiva. '
  'Imobilizar sem indicação e sem consentimento é conduta inadequada.'),
 ('Qual recurso vocês implantariam amanhã na clínica?',
  'Agenda visual impressa e vídeo gravado no celular: custo quase zero e com ensaio clínico '
  'mostrando ganho de cooperação e redução do número de consultas.'),
 ('O trabalho diz que não existe protocolo único. Então o que ele conclui?',
  'Que a conduta é individualizada, mas os princípios são constantes: anamnese ampliada, '
  'previsibilidade, suporte visual, adaptação sensorial, reforço positivo e prevenção.'),
]:
    p(q, 9.5, True, False, None, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0)
    p(a, 9.5, False, False, None, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.0, 4, recuo=0.45)

doc.save('Guia_Rapido_TEA_TDAH_TAG.docx')
print('ok')
