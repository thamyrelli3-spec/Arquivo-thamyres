# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

TITULO = ("MANEJO ODONTOLÓGICO DO PACIENTE COM TRANSTORNO DO ESPECTRO AUTISTA NÍVEL 3, "
          "TDAH E TRANSTORNO DE ANSIEDADE GENERALIZADA: REVISÃO DE LITERATURA")

CABECALHO = [
    "Autores: Thamyres Tosarelli e Beatriz Gouvea",
    "Orientadores: Rosemary Baptista Martins Teixeira e Ricardo Matsura Kodama",
]

CORPO = open('corpo.txt', encoding='utf-8').read().strip()

PALAVRAS_CHAVE = ("Palavras-chave: Transtorno do Espectro Autista; Assistência Odontológica "
                  "para Pessoas com Deficiências; Manejo Comportamental.")

REFS = [
 "AMERICAN ACADEMY OF PEDIATRIC DENTISTRY. Behavior guidance for the pediatric dental patient. "
 "In: The Reference Manual of Pediatric Dentistry. Chicago: AAPD, 2024. p. 358-378.",

 "ALBHAISI, Ismail Nabil et al. Effectiveness of psychological techniques in dental management for "
 "children with autism spectrum disorder: a systematic literature review. BMC Oral Health, Londres, "
 "v. 22, 2022. DOI: 10.1186/s12903-022-02200-7.",

 "BALIAN, Araxi et al. Is visual pedagogy effective in improving cooperation towards oral hygiene "
 "and dental care in children with autism spectrum disorder? A systematic review and meta-analysis. "
 "International Journal of Environmental Research and Public Health, Basileia, v. 18, n. 2, p. 789, 2021.",

 "BEZERRA, R. C.; ASSIS, J. A.; SANTOS, P. U. O atendimento odontológico à crianças com Transtorno "
 "do Espectro Autista: uma revisão de literatura. Brazilian Journal of Health Review, Curitiba, "
 "v. 6, n. 3, p. 13155-13171, 2023.",

 "CERMAK, Sharon A. et al. Sensory adapted dental environments to enhance oral care for children "
 "with autism spectrum disorders: a randomized controlled pilot study. Journal of Autism and "
 "Developmental Disorders, Nova York, v. 45, n. 9, p. 2876-2888, 2015.",

 "DRUMOND, Victor Zanetti et al. Dental caries in children with attention deficit/hyperactivity "
 "disorder: a meta-analysis. Caries Research, Basileia, v. 56, n. 1, p. 3-14, 2022.",

 "STEIN DUKER, Leah I. et al. Sensory adaptations to improve physiological and behavioral distress "
 "during dental visits in autistic children: a randomized crossover trial. JAMA Network Open, "
 "Chicago, v. 6, n. 6, e2316346, 2023.",
]

doc = Document()
sec = doc.sections[0]
for attr in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
    setattr(sec, attr, Cm(2.5))

def par(texto, tamanho=12, negrito=False, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY,
        espaco=1.5, depois=6):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    pf = p.paragraph_format
    pf.line_spacing = espaco
    pf.space_before = Pt(0)
    pf.space_after = Pt(depois)
    r = p.add_run(texto)
    r.bold = negrito
    r.font.name = 'Arial'
    r.font.size = Pt(tamanho)
    r._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    return p

par(TITULO, 14, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 12)
for linha in CABECALHO:
    par(linha, 14, False, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 0)
par("", 12, False, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0)
par(CORPO, 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 12)
par(PALAVRAS_CHAVE, 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 18)
par("REFERÊNCIAS BIBLIOGRÁFICAS", 14, False, WD_ALIGN_PARAGRAPH.LEFT, 1.5, 6)
for r in sorted(REFS):
    par(r, 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.0, 6)

doc.save('Resumo_Banner_PNE_TEA_TDAH_TAG.docx')
print('corpo:', len(CORPO.split()), 'palavras |', len(REFS), 'referencias')
