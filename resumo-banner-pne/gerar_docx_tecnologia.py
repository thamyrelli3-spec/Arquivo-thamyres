# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

TITULO = ("RECURSOS TECNOLÓGICOS NO ATENDIMENTO ODONTOLÓGICO DO PACIENTE COM "
          "NECESSIDADES ESPECIAIS: REVISÃO DE LITERATURA")

CABECALHO = [
    "AUTORA: THAMYRES TOSARELLI",
    "CO-AUTORA: BEATRIZ GOUVEA",
    "ORIENTADORA: PROFA. DRA. ROSEMARY MARTINS",
]

CORPO = open('corpo_tecnologia.txt', encoding='utf-8').read().strip()

PALAVRAS_CHAVE = ("Palavras-chave: Assistência Odontológica para Pessoas com Deficiências; "
                  "Tecnologia Odontológica; Teleodontologia.")

REFS = [
 "AMERICAN ACADEMY OF PEDIATRIC DENTISTRY. Behavior guidance for the pediatric dental patient. "
 "In: The Reference Manual of Pediatric Dentistry. Chicago: AAPD, 2024. p. 358-378.",

 "AMERICAN ACADEMY OF PEDIATRIC DENTISTRY. Policy on the use of silver diamine fluoride for "
 "pediatric dental patients. In: The Reference Manual of Pediatric Dentistry. Chicago: AAPD, 2024.",

 "BRASIL. Ministério da Saúde. Guia de atenção à saúde bucal da pessoa com deficiência. "
 "Brasília: Ministério da Saúde, 2019.",

 "CERMAK, Sharon A. et al. Sensory adapted dental environments to enhance oral care for children "
 "with autism spectrum disorders: a randomized controlled pilot study. Journal of Autism and "
 "Developmental Disorders, Nova York, v. 45, n. 9, p. 2876-2888, 2015.",

 "CIRIO, Silvia et al. Use of visual pedagogy to help children with ASDs facing the first dental "
 "examination: a randomized controlled trial. Children, Basileia, v. 9, n. 5, p. 729, 2022.",

 "DA SILVA MORO, Juliana et al. Efficacy of the video modeling technique as a facilitator of "
 "non-invasive dental care in autistic children: randomized clinical trial. Journal of Autism and "
 "Developmental Disorders, Nova York, v. 54, n. 2, p. 501-508, 2024.",

 "KENGNE TALLA, Pascaline et al. Teledentistry for improving access to, and quality of oral health "
 "care: overview of systematic reviews and meta-analyses. Journal of Medical Internet Research, "
 "Toronto, v. 27, e65211, 2025.",

 "LI, Ting et al. Er:YAG laser application in caries removal and cavity preparation in children: "
 "a meta-analysis. Lasers in Medical Science, Londres, v. 34, n. 2, p. 273-280, 2019.",

 "NEGI, Sapna et al. Artificial intelligence in dental caries diagnosis and detection: an umbrella "
 "review. Clinical and Experimental Dental Research, Hoboken, v. 10, n. 4, e70004, 2024.",

 "ORGANIZAÇÃO MUNDIAL DA SAÚDE. Global report on health equity for persons with disabilities. "
 "Genebra: OMS, 2022.",

 "STEIN DUKER, Leah I. et al. Sensory adaptations to improve physiological and behavioral distress "
 "during dental visits in autistic children: a randomized crossover trial. JAMA Network Open, "
 "Chicago, v. 6, n. 6, e2316346, 2023.",
]

doc = Document()
sec = doc.sections[0]
for attr in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
    setattr(sec, attr, Cm(2.5))

def par(texto, tamanho=12, negrito=False, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY,
        espaco=1.5, depois=6):
    p = doc.add_paragraph(); p.alignment = alinhamento
    pf = p.paragraph_format; pf.line_spacing = espaco
    pf.space_before = Pt(0); pf.space_after = Pt(depois)
    r = p.add_run(texto); r.bold = negrito
    r.font.name = 'Arial'; r.font.size = Pt(tamanho)
    r._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    return p

par(TITULO, 14, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 12)
for linha in CABECALHO:
    par(linha, 14, False, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 0)
par("", 12, False, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0)
par(CORPO, 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 12)
par(PALAVRAS_CHAVE, 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, 18)
par("REFERÊNCIAS BIBLIOGRÁFICAS", 14, False, WD_ALIGN_PARAGRAPH.LEFT, 1.5, 6)
for r in REFS:
    par(r, 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.0, 6)

doc.save('Resumo_Banner_PNE_Tecnologia.docx')
print('corpo:', len(CORPO.split()), 'palavras |', len(REFS), 'referencias')
