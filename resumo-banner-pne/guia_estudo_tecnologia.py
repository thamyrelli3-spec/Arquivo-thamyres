# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
sec = doc.sections[0]
for a in ('top_margin','bottom_margin','left_margin','right_margin'):
    setattr(sec, a, Cm(2.0))

def p(texto='', tam=11, neg=False, al=WD_ALIGN_PARAGRAPH.JUSTIFY, esp=1.15, dep=6, recuo=0):
    par = doc.add_paragraph(); par.alignment = al
    pf = par.paragraph_format; pf.line_spacing = esp
    pf.space_before = Pt(0); pf.space_after = Pt(dep)
    if recuo: pf.left_indent = Cm(recuo)
    r = par.add_run(texto); r.bold = neg
    r.font.name = 'Arial'; r.font.size = Pt(tam)
    r._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    return par

def h(t, tam=13): p(t, tam, True, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 8)
def b(t, tam=11): p('•  ' + t, tam, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 4, recuo=0.6)

p('GUIA DE ESTUDO — BANNER DE PNE', 14, True, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 4)
p('Recursos tecnológicos no atendimento do paciente com necessidades especiais',
  12, False, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 16)

h('1. DE ONDE SAIU CADA PARTE DO RESUMO')
mapa = [
 ('"barreiras de acesso e pior saúde bucal"',
  'OMS (2022) e BRASIL, Ministério da Saúde (2019).'),
 ('"pedagogia visual digital prepara o paciente antes da consulta"',
  'CIRIO et al. (2022) — ensaio clínico randomizado com vídeo e fotos antes do primeiro exame.'),
 ('"modelagem por vídeo reduz o número de sessões"',
  'DA SILVA MORO et al. (2024) — ensaio clínico randomizado; é o dado mais concreto do resumo.'),
 ('"escova elétrica quando a higiene depende do cuidador"',
  'BRASIL (2019); há também ensaio clínico comparando escova manual e sônica em pessoas com '
  'deficiência intelectual, citado no seu seminário.'),
 ('"ambiente sensorialmente adaptado reduz o estresse fisiológico e comportamental"',
  'CERMAK et al. (2015) e STEIN DUKER et al. (2023).'),
 ('"teleodontologia amplia o acesso e orienta o cuidador"',
  'KENGNE TALLA et al. (2025) — overview de revisões sistemáticas.'),
 ('"diagnóstico auxiliado por inteligência artificial"',
  'NEGI et al. (2024) — revisão guarda-chuva sobre IA na detecção de cárie.'),
 ('"laser para remoção de cárie sem o ruído da alta rotação"',
  'LI et al. (2019) — metanálise de Er:YAG em crianças.'),
 ('"monitorização durante a sedação"',
  'AAPD (2024) — o Reference Manual traz a diretriz de sedação e monitorização.'),
 ('"a tecnologia não substitui o vínculo e o manejo comportamental"',
  'AAPD (2024) — a tecnologia entra como ferramenta dentro do manejo básico, não no lugar dele.'),
]
for frase, fonte in mapa:
    p(frase, 11, True, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 0)
    p(fonte, 11, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 8, recuo=0.6)

doc.add_page_break()
h('2. RESUMO DAS REFERÊNCIAS')
refs = [
 ('OMS (2022) — Global report on health equity for persons with disabilities',
  'Relatório da Organização Mundial da Saúde, Genebra.',
  ['Documenta a iniquidade em saúde das pessoas com deficiência e as barreiras de acesso aos '
   'serviços.',
   'Serve para abrir o banner: justifica por que o tema importa antes de falar de tecnologia.']),

 ('BRASIL, Ministério da Saúde (2019) — Guia de atenção à saúde bucal da pessoa com deficiência',
  'Brasília, publicação oficial, acesso livre.',
  ['Orienta os profissionais da rede SUS sobre manejo e cuidado, com foco nas condições de maior '
   'demanda no consultório: deficiência intelectual, demência, Parkinson, paralisia cerebral, '
   'síndrome de Down e TEA.',
   'É a referência nacional do trabalho. Leia esta primeiro, é em português e gratuita.']),

 ('CIRIO et al. (2022) — Pedagogia visual antes do primeiro exame',
  'Children, v. 9, n. 5, p. 729. Ensaio clínico randomizado.',
  ['113 crianças com TEA sorteadas para dois grupos de preparo: vídeo ou fotos. 84 concluíram '
   '(vídeo n=41, fotos n=43).',
   'Avaliou quantas das 8 etapas do exame a criança alcançou e o nível de cooperação pela escala '
   'de Frankl.',
   'Conclusão: preparar a criança antes da consulta funciona. Vale saber o que é a escala de '
   'Frankl, é pergunta provável.']),

 ('DA SILVA MORO et al. (2024) — Modelagem por vídeo',
  'J Autism Dev Disord, v. 54, n. 2, p. 501-508. Ensaio clínico randomizado cego. Autores '
  'brasileiros, da UFSC.',
  ['Mostrou que a modelagem por vídeo reduz o número de consultas necessárias para o atendimento '
   'não invasivo em crianças autistas.',
   'É o melhor dado do seu resumo: tecnologia barata, resultado clínico objetivo e pesquisa '
   'nacional. Bom argumento na apresentação.']),

 ('CERMAK et al. (2015) e STEIN DUKER et al. (2023) — Ambiente sensorialmente adaptado',
  'J Autism Dev Disord, v. 45, n. 9, p. 2876-2888; e JAMA Network Open, v. 6, n. 6, e2316346.',
  ['Cermak: piloto com 44 crianças, cada uma atendida em sala convencional e em sala adaptada, em '
   'ordem sorteada. Adaptações: luz reduzida, projeção no teto, música calma e colete com peso.',
   'Stein Duker: mesma linha, agora com 162 crianças, ensaio cruzado. Estresse fisiológico '
   'significativamente menor no ambiente adaptado, em todas as fases da consulta.',
   'É a evidência mais forte do trabalho. Se perguntarem qual estudo tem melhor qualidade '
   'metodológica, é o de 2023.']),

 ('KENGNE TALLA et al. (2025) — Teleodontologia',
  'J Med Internet Res, v. 27, e65211. Overview de revisões sistemáticas e metanálises.',
  ['Conclui que a teleodontologia é alternativa efetiva e eficiente ao atendimento presencial.',
   'Ressalva importante: a qualidade das revisões incluídas é preocupante, faltam estudos '
   'metodologicamente rigorosos. Cite a ressalva, mostra leitura crítica.',
   'No Brasil, a Resolução CFO nº 226/2020 é o que regulamenta a odontologia a distância. Vale '
   'saber que o teleatendimento não substitui a consulta presencial de diagnóstico.']),

 ('NEGI et al. (2024) — Inteligência artificial no diagnóstico de cárie',
  'Clin Exp Dent Res, v. 10, n. 4, e70004. Revisão guarda-chuva (umbrella review).',
  ['Reúne 7 revisões sistemáticas sobre modelos de IA na detecção e no diagnóstico de cárie.',
   'Saiba explicar o que é uma revisão guarda-chuva: é a revisão que reúne outras revisões '
   'sistemáticas, fica no topo da pirâmide de evidência.',
   'Aplicação no PNE: ajuda quando a radiografia sai de qualidade ruim porque o paciente não '
   'colabora.']),

 ('LI et al. (2019) — Laser Er:YAG em crianças',
  'Lasers in Medical Science, v. 34, n. 2, p. 273-280. Metanálise de 7 ensaios randomizados.',
  ['Avalia o laser Er:YAG na remoção de cárie e no preparo cavitário em crianças.',
   'Relevância para o PNE: remove tecido sem o ruído e a vibração da alta rotação, dois dos '
   'estímulos que mais desencadeiam recusa em pacientes com hipersensibilidade sensorial.',
   'ATENÇÃO: no seu seminário essa referência está com a revista errada (aparece como European '
   'Journal of Paediatric Dentistry, 2018). O correto é Lasers in Medical Science, 2019, v. 34, '
   'n. 2, p. 273-280, autores LI, Ting et al. Corrija também no deck.']),

 ('AAPD (2024) — Behavior guidance for the pediatric dental patient',
  'Reference Manual of Pediatric Dentistry, p. 358-378.',
  ['Diretriz que separa manejo básico (comunicação, dizer-mostrar-fazer, reforço positivo, '
   'distração, dessensibilização) de manejo avançado (contenção protetora, sedação e anestesia '
   'geral).',
   'Use para defender a tese do banner: a tecnologia é ferramenta dentro do manejo básico. Ela '
   'adia ou evita o avançado, não substitui o vínculo com o paciente.']),
]
for titulo, dados, itens in refs:
    p(titulo, 11.5, True, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 0)
    p(dados, 10.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 4)
    for i in itens: b(i)
    p('', 8, dep=6)

doc.add_page_break()
h('3. O QUE ESTUDAR ALÉM DAS REFERÊNCIAS')
for t in [
 'Quem é o paciente com necessidades especiais: a definição não é só deficiência, inclui condições '
 'sistêmicas, transtornos do neurodesenvolvimento, condições temporárias e extremos de idade.',
 'Lei 13.146/2015, a Lei Brasileira de Inclusão: o conceito de barreira (urbanística, '
 'arquitetônica, de transporte, de comunicação, atitudinal e tecnológica). O banner fica muito bem '
 'amarrado se você mostrar a tecnologia derrubando cada barreira.',
 'ABNT NBR 9050:2020: acessibilidade de edificações. Serve para a parte de acessibilidade física '
 'do consultório.',
 'Resolução CFO nº 226/2020: o que a teleodontologia pode e não pode fazer no Brasil.',
 'Censo 2022 do IBGE: dados de pessoas com deficiência no país, bom número de abertura.',
 'Escala de Frankl: como se classifica o comportamento da criança (de definitivamente negativo a '
 'definitivamente positivo). Aparece no estudo do Cirio.',
 'Diferença entre modelagem por vídeo e pedagogia visual: a primeira mostra alguém passando pelo '
 'procedimento, a segunda usa figuras em sequência para antecipar as etapas.',
 'Realidade virtual como distração: saiba dizer que a evidência ainda é menor que a do ambiente '
 'adaptado, e que há limitação em pacientes com hipersensibilidade, que podem não tolerar os óculos.',
 'Diamino fluoreto de prata e ART: tecnologia de baixo custo que evita o preparo cavitário. '
 'Combina com a tese do trabalho.',
 'Pirâmide de evidência: onde entram ensaio clínico randomizado, revisão sistemática, metanálise, '
 'revisão guarda-chuva e diretriz. Você cita os cinco tipos no banner.',
]:
    b(t)

p('', 8, dep=10)
h('4. PERGUNTAS QUE PODEM TE FAZER')
qa = [
 ('Qual desses recursos você implantaria hoje na clínica-escola?',
  'Modelagem por vídeo e pedagogia visual. Custam quase nada, funcionam com o celular e têm ensaio '
  'clínico mostrando redução do número de consultas. Depois, adaptações sensoriais simples: '
  'diminuir a luz, óculos escuros, abafador de ruído.'),
 ('Tecnologia não é desculpa para deixar de aprender manejo?',
  'Não. A diretriz da AAPD trata esses recursos como ferramentas dentro do manejo básico. Sem '
  'vínculo, comunicação e previsibilidade, nenhum equipamento resolve.'),
 ('Tudo isso não é caro demais para o SUS?',
  'Os recursos com melhor evidência são justamente os mais baratos. O que é caro, como a realidade '
  'virtual e a inteligência artificial, é o que tem evidência menos consolidada.'),
 ('Qual a limitação da teleodontologia?',
  'A própria revisão do Kengne Talla aponta que a qualidade metodológica das revisões é baixa. '
  'E no Brasil a Resolução CFO 226/2020 delimita o que pode ser feito a distância: não substitui o '
  'exame clínico presencial.'),
 ('Por que o laser ajudaria um paciente com necessidades especiais?',
  'Porque elimina o ruído e a vibração da alta rotação, que são dois dos gatilhos sensoriais mais '
  'comuns de recusa. A metanálise do Li mostra a eficácia na remoção de cárie em crianças.'),
]
for q, a in qa:
    p(q, 11, True, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 0)
    p(a, 11, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 8, recuo=0.6)

p('', 8, dep=10)
h('5. ONDE LER CADA UMA')
for l in [
 'BRASIL (2019): bvsms.saude.gov.br/bvs/publicacoes/guia_atencao_saude_bucal_pessoa_deficiencia.pdf',
 'OMS (2022): who.int — Global report on health equity for persons with disabilities',
 'CIRIO et al. (2022): pmc.ncbi.nlm.nih.gov/articles/PMC9139454',
 'DA SILVA MORO et al. (2024): pubmed.ncbi.nlm.nih.gov/36357551',
 'CERMAK et al. (2015): pubmed.ncbi.nlm.nih.gov/25931290',
 'STEIN DUKER et al. (2023): ncbi.nlm.nih.gov/pmc/articles/PMC10238943',
 'KENGNE TALLA et al. (2025): jmir.org/2025/1/e65211',
 'NEGI et al. (2024): pmc.ncbi.nlm.nih.gov/articles/PMC11358700',
 'LI et al. (2019): pubmed.ncbi.nlm.nih.gov/30003427',
 'AAPD (2024): aapd.org/globalassets/media/policies_guidelines/bp_behavguide.pdf',
]:
    b(l, 10.5)

doc.save('Guia_de_estudo_Tecnologia.docx')
print('ok')
